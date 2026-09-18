from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from src.attribution import (append_record, build_input, config_hash, digest, load_journal,
                             offline_model, run_attribution, summarize, tool_trace)
from src.attribution_report import anchor, generate_reports
from src.attribution_schema import STAGES, parse_result
from src.config import AttributionConfig, AttributionConfigLoader, ModelConfig, SanitizeRule
from src.cli import _run_both
from src.dataset import load_dataset
from src.importer import import_data


def config(**overrides):
    return AttributionConfig(model=ModelConfig(provider="test", model="fake", base_url="https://example.invalid"),
                             prompt="只使用提供的材料。", retries=0, **overrides)


def source_row():
    return {"row_index": 0, "query": "谁可以申请？", "human_note": "回答遗漏了首次申请的条件。",
            "rendered_request": {"messages": [
                {"role": "user", "content": "谁可以申请？"},
                {"role": "assistant", "tool_calls": [{"id": "call_01", "function": {
                    "name": "WebSearch", "arguments": '{"queries":["申请条件"]}'}}]},
                {"role": "tool", "tool_call_id": "call_01", "content": "初回申請者のみ対象です。"}]},
            "response": {"choices": [{"message": {"content": "所有人都可以申请。"}}]}}


def good_result():
    return {
        "status": "primary_identified", "content_summaries_zh": {"query": None, "answer": None, "human_note": None},
        "note_assessment": [{"claim": "遗漏了首次申请的条件", "assessment": "supported", "reason": "原文有限定。", "evidence_ids": ["E1", "E2"]}],
        "stage_checks": [{"stage": stage, "status": "checked", "finding": "检查了已提供材料。", "evidence_ids": []} for stage in STAGES],
        "issues": [{"id": "I1", "description": "答复扩大适用范围。", "impact_level": "core", "evidence_ids": ["E1", "E2"]}],
        "evidence": [
            {"id": "E1", "source": "messages", "path": "/messages/2/content", "quote": "初回申請者のみ対象です。",
             "source_language": "ja", "translation_zh": "仅适用于首次申请者。", "translation_note": None},
            {"id": "E2", "source": "answer", "path": "/answer", "quote": "所有人都可以申请。",
             "source_language": "zh", "translation_zh": None, "translation_note": None}],
        "factors": [{"id": "F1", "stage": "answer", "problem_type": "限定条件丢失", "description": "删除了首次申请限定。",
            "issue_ids": ["I1"], "priority": 1, "impact_level": "core", "evidence_strength": "direct", "evidence_ids": ["E1", "E2"],
            "mechanism": "来源有限定，答复未保留。", "ranking_reason": "直接影响核心结论。",
            "fix": {"action": "保留适用范围。", "validation": "用限定条件案例复测。"}}],
        "primary_factor_ids": ["F1"], "suggestions": [], "open_questions": [], "summary": "答复遗漏限定，优先改善证据使用。"}


def model_response(result=None):
    return ({"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(result or good_result(), ensure_ascii=False)}}]}, {})


class AttributionSchemaTests(unittest.TestCase):
    def setUp(self):
        self.snapshot = build_input(source_row(), config())[0]

    def test_multilingual_evidence_and_arbitrary_factors(self):
        result = good_result()
        for i in range(2, 9):
            factor = deepcopy(result["factors"][0])
            factor.update(id=f"F{i}", priority=i)
            result["factors"].append(factor)
        parsed = parse_result(json.dumps(result), self.snapshot)
        self.assertEqual(len(parsed["factors"]), 8)
        self.assertEqual(parsed["evidence"][0]["quote"], self.snapshot["messages"][2]["content"])

    def test_invalid_evidence_relations_and_states(self):
        mutations = [
            lambda r: r["evidence"][0].update(quote="仅适用于首次申请者。"),
            lambda r: r["evidence"][0].update(translation_zh=None),
            lambda r: r["evidence"][0].update(path="/messages/99/content"),
            lambda r: r["evidence"][0].update(source="query"),
            lambda r: r["factors"][0].update(issue_ids=["missing"]),
            lambda r: r["factors"][0].update(priority=True),
            lambda r: r.update(status="joint_primary"),
            lambda r: r.update(status="no_supported_issue", primary_factor_ids=[]),
            lambda r: r["stage_checks"].pop(),
            lambda r: r["evidence"].append(deepcopy(r["evidence"][0])),
        ]
        for mutation in mutations:
            with self.subTest(mutation=mutation):
                result = good_result()
                mutation(result)
                with self.assertRaises(ValueError):
                    parse_result(json.dumps(result), self.snapshot)

    def test_note_is_not_sufficient_evidence(self):
        result = good_result()
        result["evidence"] = [{"id": "E1", "source": "human_note", "path": "/human_note",
            "quote": self.snapshot["human_note"], "source_language": "zh", "translation_zh": None, "translation_note": None}]
        for item in [*result["issues"], *result["factors"], *result["note_assessment"]]:
            item["evidence_ids"] = ["E1"]
        with self.assertRaisesRegex(ValueError, "人工评论"):
            parse_result(json.dumps(result), self.snapshot)

    def test_joint_unresolved_insufficient_and_no_issue(self):
        result = good_result()
        second = deepcopy(result["factors"][0])
        second["id"] = "F2"
        result["factors"].append(second)
        result.update(status="joint_primary", primary_factor_ids=["F1", "F2"])
        parse_result(json.dumps(result), self.snapshot)
        result.update(status="ranking_unresolved", primary_factor_ids=[])
        for factor in result["factors"]:
            factor["priority"] = None
        result["open_questions"] = [{"question": "哪项影响更大？", "needed_material": "修复对照", "limits": "无法比较影响"}]
        parse_result(json.dumps(result), self.snapshot)
        result.update(status="insufficient_evidence", factors=[])
        parse_result(json.dumps(result), self.snapshot)
        result.update(status="no_supported_issue", issues=[])
        parse_result(json.dumps(result), self.snapshot)

    def test_search_matching_and_recursive_sanitization(self):
        row = source_row()
        row["query"] = "SECRET"
        row["human_note"] = "SECRET"
        row["rendered_request"]["messages"][1]["tool_calls"][0]["function"]["arguments"] = {"queries": ["SECRET"]}
        row["rendered_request"]["messages"][2]["content"] = {"text": "SECRET"}
        cfg = config(sanitize=[SanitizeRule(pattern="SECRET", replacement="[MASK]")])
        snapshot, _, trace = build_input(row, cfg)
        self.assertNotIn("SECRET", json.dumps(snapshot))
        self.assertTrue(trace["calls"][0]["is_search"])
        self.assertTrue(trace["results"][0]["is_search"])
        row["rendered_request"]["messages"][1]["tool_calls"][0]["function"]["name"] = "calculator"
        self.assertFalse(build_input(row, cfg)[2]["calls"][0]["is_search"])
        self.assertEqual(row["query"], "SECRET")

    def test_config_hash_and_offline_validation(self):
        cfg = config()
        changed = cfg.model_copy(deep=True)
        changed.concurrency = 4
        self.assertEqual(config_hash(cfg), config_hash(changed))
        changed.prompt += "new rule"
        self.assertNotEqual(config_hash(cfg), config_hash(changed))
        cfg.model.extra_body = {"enable_search": True}
        with self.assertRaises(ValueError):
            offline_model(cfg)
        self.assertFalse(AttributionConfigLoader().get_attribution().model.extra_body["enable_search"])


class AttributionIntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def test_combined_config_failure_does_not_block_attribution(self):
        attribution = AsyncMock(return_value=Path("output"))
        with patch("src.cli.EvalConfigLoader", side_effect=ValueError("bad judge config")), \
             patch("src.cli.AttributionConfigLoader") as loader, patch("src.cli.run_attribution", attribution):
            loader.return_value.get_attribution.return_value = config()
            outcomes = await _run_both("sample", 1, False)
            self.assertIsInstance(outcomes[0], ValueError)
            attribution.assert_awaited_once()

    async def test_combined_attribution_failure_does_not_block_judge(self):
        evaluation = AsyncMock()
        with patch("src.cli.EvalConfigLoader") as loader, \
             patch("src.cli.AttributionConfigLoader", side_effect=ValueError("bad attribution config")), \
             patch("src.cli.run_evaluation", evaluation):
            loader.return_value.get_eval.return_value = object()
            outcomes = await _run_both("sample", 1, False)
            self.assertIsInstance(outcomes[1], ValueError)
            evaluation.assert_awaited_once()

    async def test_independent_resume_input_invalidation_and_report_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run = root / "sample"
            run.mkdir()
            row = source_row()
            source = run / "responses.jsonl"
            source.write_text(json.dumps(row, ensure_ascii=False) + "\n", encoding="utf-8")
            # 评分文件故意不是合法 JSON；归因不能读取它。
            (run / "scores.jsonl").write_text("DO NOT READ OR MODIFY", encoding="utf-8")
            client = AsyncMock()
            api = AsyncMock(return_value=model_response())
            with patch("src.attribution.RESULTS_DIR", root), patch("src.attribution.create_client", return_value=client), patch("src.attribution.call_model", api):
                directory = await run_attribution("sample", config(), formats=("html", "xlsx"))
                await run_attribution("sample", config(), formats=())
                self.assertEqual(api.await_count, 1)
                client.close.assert_awaited_once()
                row["human_note"] += "请复核。"
                source.write_text(json.dumps(row, ensure_ascii=False) + "\n", encoding="utf-8")
                await run_attribution("sample", config(), report_only=True, formats=("html",))
                summary = json.loads((directory / "attribution_summary.json").read_text(encoding="utf-8"))
                self.assertEqual(summary["execution_counts"], {"not_processed": 1})
                self.assertEqual(api.await_count, 1)
                await run_attribution("sample", config(), formats=())
                self.assertEqual(api.await_count, 2)
                self.assertEqual(len(load_journal(directory / "attributions.jsonl")), 2)
                await run_attribution("sample", config(), force=True, formats=())
                self.assertEqual(api.await_count, 3)
            self.assertEqual((run / "scores.jsonl").read_text(), "DO NOT READ OR MODIFY")

    async def test_only_with_human_note_filters_effective_notes(self):
        from src.human_notes import list_cases, save_notes
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run = root / "r"
            run.mkdir()
            rows = [source_row() for _ in range(7)]
            for index, row in enumerate(rows):
                row["row_index"] = index
                row.pop("human_note", None)
            rows[0]["human_note"] = "人工评论"
            rows[1]["human_note"] = " \n\t "
            rows[2]["note"] = "兼容评论"
            rows[4]["note"] = "将被清空"
            source = run / "responses.jsonl"
            def write_source():
                source.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
            write_source()
            states = list_cases(run)
            save_notes(run, [{"row_index": i, "human_note": note, "revision": states[i]["revision"]}
                             for i, note in [(3, "页面评论"), (4, ""), (5, "旧版评论")]])
            rows[5]["query"] += "已修改"
            write_source()
            api = AsyncMock(return_value=model_response())
            with patch("src.attribution.RESULTS_DIR", root), patch("src.attribution.create_client", return_value=AsyncMock()), patch("src.attribution.call_model", api):
                directory = await run_attribution("r", config(), only_with_human_note=True, formats=())
                self.assertEqual({r["row_index"] for r in load_journal(directory / "attributions.jsonl")}, {0, 2, 3})
                self.assertEqual(api.await_count, 3)
                await run_attribution("r", config(), only_with_human_note=True, formats=())
                self.assertEqual(api.await_count, 3)
                await run_attribution("r", config(), only_with_human_note=True, rows=[1, 2], force=True, formats=())
                self.assertEqual(api.await_count, 4)
                self.assertEqual(load_journal(directory / "attributions.jsonl")[-1]["row_index"], 2)
                await run_attribution("r", config(), only_with_human_note=True, rows=[1, 4, 5, 6], force=True, formats=())
                self.assertEqual(api.await_count, 4)
                await run_attribution("r", config(), only_with_human_note=True, report_only=True, formats=())
                self.assertEqual(api.await_count, 4)
                stats = json.loads((directory / "attribution_summary.json").read_text(encoding="utf-8"))
                self.assertEqual(stats["execution_counts"], {"success": 3, "not_processed": 4})
                await run_attribution("r", config(), formats=())
                self.assertEqual(api.await_count, 8)

    async def test_failure_retries_and_oversize_without_truncation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "r").mkdir()
            row = source_row()
            row["query"] = "超长" * 30000
            (root / "r" / "responses.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
            api = AsyncMock(return_value=model_response())
            with patch("src.attribution.RESULTS_DIR", root), patch("src.attribution.create_client", return_value=AsyncMock()) as factory, patch("src.attribution.call_model", api):
                directory = await run_attribution("r", config(max_input_chars=100), formats=())
                factory.assert_not_called()
                record = load_journal(directory / "attributions.jsonl")[0]
                self.assertEqual(record["input_snapshot"]["query"], row["query"])
                self.assertEqual(record["execution"]["status"], "failed")
                api.side_effect = [ValueError("temporary failure"), model_response()]
                cfg = config()
                cfg.retries = 1
                cfg.retry_delay = 0
                directory = await run_attribution("r", cfg, formats=())
                self.assertEqual(load_journal(directory / "attributions.jsonl")[-1]["execution"]["status"], "success")
                self.assertEqual(api.await_count, 2)

    async def test_streaming_selected_rows_and_partial_completion(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "r").mkdir()
            rows = [source_row(), source_row()]
            rows[1]["row_index"] = 1
            (root / "r" / "responses.jsonl").write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")
            cfg = config()
            cfg.model.stream = True
            response, request = model_response()
            stream = AsyncMock(return_value=(response, request, 0.1))
            with patch("src.attribution.RESULTS_DIR", root), patch("src.attribution.create_client", return_value=AsyncMock()), patch("src.attribution.call_model_stream", stream):
                directory = await run_attribution("r", cfg, rows=[1], formats=())
                self.assertEqual(stream.await_count, 1)
                stats = json.loads((directory / "attribution_summary.json").read_text(encoding="utf-8"))
                self.assertEqual(stats["execution_counts"], {"not_processed": 1, "success": 1})
                with self.assertRaises(ValueError):
                    await run_attribution("r", cfg, rows=[99], formats=())


class AttributionStorageReportTests(unittest.TestCase):
    def test_invalid_tool_ids_are_reported_as_missing_links(self):
        messages = [{"role": "assistant", "tool_calls": [{"id": {}, "function": {"name": "WebSearch"}}]},
                    {"role": "tool", "tool_call_id": [], "content": "text"}]
        trace = tool_trace(messages, ["WebSearch"])
        self.assertEqual(len(trace["warnings"]), 2)
        self.assertIsNone(trace["results"][0]["is_search"])

    def test_partial_journal_recovery_preserves_previous_records(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "attributions.jsonl"
            path.write_bytes(b'{"case_id":"old"}\n{"incomplete":"\xe4')
            self.assertEqual(len(load_journal(path)), 1)
            append_record(path, {"case_id": "new"})
            self.assertEqual([r["case_id"] for r in load_journal(path)], ["old", "new"])
            self.assertEqual(len(list(Path(tmp).glob("*.partial"))), 1)

    def test_reports_full_text_safe_html_excel_and_multi_factor_statistics(self):
        snapshot, quality, trace = build_input(source_row(), config())
        malicious = '</script><img src=x onerror="alert(1)">'
        snapshot["query"] = "=1+1" + "长" * 40000 + malicious
        result = good_result()
        second = deepcopy(result["factors"][0])
        second.update(id="F2", priority=2)
        result["factors"].append(second)
        record = {"case_id": "r:0", "row_index": 0, "input_hash": "h", "config_hash": "c", "input_snapshot": snapshot,
                  "input_quality": quality, "tool_trace": trace, "execution": {"status": "success"}, "result": result}
        summary = summarize([record])
        self.assertEqual(summary["involved_cases_by_stage"], {"answer": 1})
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            generate_reports(directory, [record], summary)
            html = (directory / "report.html").read_text(encoding="utf-8")
            self.assertNotIn(malicious, html)
            self.assertIn("&lt;img", html)
            self.assertIn("长" * 40000, html)
            self.assertIn('id="' + anchor("r:0", "/messages/2/content") + '"', html)
            self.assertIn("初回申請者のみ対象です。", html)
            book = load_workbook(directory / "report.xlsx")
            self.assertIn("展示摘录", book["案例总览"]["C2"].value)
            self.assertEqual(book["案例总览"]["C2"].data_type, "s")
            self.assertEqual(book["因素明细"].max_row, 3)
            self.assertEqual(book["因素明细"]["C2"].data_type, "n")
            self.assertEqual(book["证据明细"]["E2"].value, "初回申請者のみ対象です。")
            self.assertEqual(book["证据明细"]["G2"].value, "仅适用于首次申请者。")
            book.close()

    def test_import_preserves_custom_note_and_answer_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.jsonl"
            source.write_text(json.dumps({"query": "问题", "response": "回答", "api_json": [], "评论": "意见", "location": "上海"}, ensure_ascii=False) + "\n", encoding="utf-8")
            with patch("src.importer.RESULTS_DIR", root / "results"):
                import_data(str(source), "r", note_col="评论")
            imported = json.loads((root / "results/r/responses.jsonl").read_text(encoding="utf-8"))
            self.assertEqual(imported["human_note"], "意见")
            self.assertEqual(imported["human_note_answer_hash"], digest("回答"))
            saved = json.loads((root / "results/r/input.jsonl").read_text(encoding="utf-8"))
            self.assertEqual(saved["评论"], "意见")
            self.assertEqual(saved["location"], "上海")

    def test_dataset_preserves_notes_on_later_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "input.jsonl"
            source.write_text(json.dumps({"query": "q", "api_json": []}) + "\n" + json.dumps({"query": "q2", "api_json": [], "note": "review"}), encoding="utf-8")
            self.assertEqual(load_dataset(str(source))[1]["note"], "review")


if __name__ == "__main__":
    unittest.main()
