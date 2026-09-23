import json
from pathlib import Path
import tempfile
import unittest
import re
import shutil
import subprocess
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from src.attribution import load_journal, run_attribution
from src.attribution_scoring import score_cases
from src.config import EvalConfig
from src.evaluator import compute_eval_hash, score_input_hash, build_score_messages, run_evaluation
from test_attribution import config, source_row, model_response


def evaluation():
    return EvalConfig(model=config().model, prompt="Judge only", dimensions=["overall", "factuality"])


def response(overall=2):
    result = {"overall": overall, "factuality": 3, "analysis": "<script>不是指令</script>", "answer_trace": None}
    return {"choices": [{"message": {"content": json.dumps(result)}}]}, {}


class ScoringTests(unittest.IsolatedAsyncioTestCase):
    async def test_real_eval_record_reuse(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run = root / "r"
            run.mkdir()
            row = source_row()
            (run / "responses.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
            cfg = evaluation()
            with patch("src.evaluator.RESULTS_DIR", root), patch("src.evaluator.create_client", return_value=AsyncMock()), patch("src.evaluator.call_model", AsyncMock(return_value=response())):
                await run_evaluation("r", cfg)
            original = (run / "scores.jsonl").read_bytes()
            with patch("src.attribution_scoring.create_client") as factory:
                scored = await score_cases(run, [row], {0}, config=cfg)
                self.assertEqual(scored[0]["source"], "eval")
                self.assertEqual(scored[0]["overall"], 2)
                factory.assert_not_called()
            self.assertEqual((run / "scores.jsonl").read_bytes(), original)

    async def test_resume_invalidation_selection_and_eval_reuse(self):
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp)
            row = source_row()
            other = dict(source_row(), row_index=1)
            cfg = evaluation()
            api = AsyncMock(return_value=response())
            with patch("src.attribution_scoring.create_client", return_value=AsyncMock()), patch("src.attribution_scoring.call_model", api):
                scored = await score_cases(run, [row, other], {0}, config=cfg)
                self.assertEqual(scored[0]["overall"], 2)
                self.assertEqual(scored[1]["status"], "not_processed")
                row["human_note"] = "NEW HUMAN NOTE"
                await score_cases(run, [row], {0}, config=cfg)
                self.assertEqual(api.await_count, 1)
                self.assertNotIn("NEW HUMAN NOTE", json.dumps(build_score_messages(row, cfg)))
                row["response"]["choices"][0]["message"]["content"] += "changed"
                scored = await score_cases(run, [row], {0}, config=cfg, report_only=True)
                self.assertEqual(scored[0]["status"], "not_processed")
                self.assertEqual(api.await_count, 1)
                await score_cases(run, [row], {0}, config=cfg)
                self.assertEqual(api.await_count, 2)
                cfg.prompt += "new standard"
                await score_cases(run, [row], {0}, config=cfg)
                self.assertEqual(api.await_count, 3)
                raw = json.loads(response()[0]["choices"][0]["message"]["content"])
                raw.update(row_index=1, eval_config_hash=compute_eval_hash(cfg), score_input_hash=score_input_hash(other))
                (run / "scores.jsonl").write_text(json.dumps(raw) + "\n", encoding="utf-8")
                scored = await score_cases(run, [other], {1}, config=cfg)
                self.assertEqual(scored[1]["source"], "eval")
                self.assertEqual(api.await_count, 3)
                # 无输入版本的旧评分不允许仅按行号套用。
                raw.pop("score_input_hash")
                raw["row_index"] = 2
                (run / "scores.jsonl").write_text(json.dumps(raw) + "\n", encoding="utf-8")
                await score_cases(run, [dict(other, row_index=2)], {2}, config=cfg)
                self.assertEqual(api.await_count, 4)

    async def test_failure_retry_empty_null_and_report_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            run = Path(tmp)
            api = AsyncMock(side_effect=RuntimeError("secret provider error"))
            with patch("src.attribution_scoring.create_client", return_value=AsyncMock()), patch("src.attribution_scoring.call_model", api), patch("src.attribution_scoring.asyncio.sleep", new_callable=AsyncMock):
                scored = await score_cases(run, [source_row()], {0}, config=evaluation())
                self.assertEqual(api.await_count, 4)
                self.assertEqual(scored[0]["status"], "failed")
                self.assertNotIn("secret", scored[0]["error"])
                await score_cases(run, [source_row()], {0}, config=evaluation(), report_only=True)
                self.assertEqual(api.await_count, 4)
                api.side_effect = None
                api.return_value = response(None)
                scored = await score_cases(run, [source_row()], {0}, config=evaluation())
                self.assertEqual(api.await_count, 5)
                self.assertEqual(scored[0]["status"], "unscorable")
                await score_cases(run, [source_row()], {0}, config=evaluation())
                self.assertEqual(api.await_count, 5)
                empty = source_row()
                empty["response"]["choices"][0]["message"]["content"] = ""
                scored = await score_cases(run, [empty], {0}, config=evaluation())
                self.assertEqual(scored[0]["status"], "unscorable")
                self.assertIsNone(scored[0]["overall"])
                self.assertEqual(api.await_count, 5)

    async def test_independent_execution_reports_and_force(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run = root / "r"
            run.mkdir()
            row = source_row()
            source = run / "responses.jsonl"
            source.write_text(json.dumps(row) + "\n", encoding="utf-8")
            score_api, attribute_api = AsyncMock(return_value=response()), AsyncMock(return_value=model_response())
            with patch("src.attribution.RESULTS_DIR", root), patch("src.attribution_scoring.EvalConfigLoader") as loader, patch("src.attribution_scoring.create_client", return_value=AsyncMock()), patch("src.attribution_scoring.call_model", score_api), patch("src.attribution.create_client", return_value=AsyncMock()), patch("src.attribution.call_model", attribute_api):
                loader.return_value.get_eval.return_value = evaluation()
                directory = await run_attribution("r", config(), with_score=True)
                self.assertEqual((score_api.await_count, attribute_api.await_count), (1, 1))
                html = (directory / "report.html").read_text(encoding="utf-8")
                self.assertIn("score-min", html)
                self.assertIn("&lt;script&gt;", html)
                workbook = load_workbook(directory / "report.xlsx")
                self.assertIn("评分明细", workbook.sheetnames)
                self.assertEqual(workbook["案例总览"].cell(2, 12).value, 2)
                workbook.close()
                node = shutil.which("node") or "C:/Program Files/nodejs/node.exe"
                if Path(node).exists():
                    script = re.search(r"<script>(.*?)</script>", html, re.S).group(1)
                    harness = """
const fields={};for(const id of ['search','stage','status','type','score-min','score-max','score-status','score-sort','count','score-stages'])fields[id]={value:'',addEventListener(){}};
const order=[];
const nodes=['2','5',''].map((score,i)=>({id:i,dataset:{score,scoreStatus:score?'success':'failed',primary:'answer|answer|',stage:'answer',status:'success',type:''},textContent:'sample',parentElement:{appendChild(c){const old=order.indexOf(c.id);if(old>=0)order.splice(old,1);order.push(c.id)}}}));
global.document={querySelectorAll:()=>nodes,getElementById:id=>fields[id]};global.window={addEventListener(){}};global.location={hash:''};
"""
                    checks = """
fields['score-max'].value='3';filter();if(nodes.filter(c=>!c.hidden).length!==1)throw Error('range/null filter');
if(!fields['score-stages'].textContent.includes('answer 1'))throw Error('deduplicate stages');
fields['score-max'].value='';fields['score-sort'].value='desc';filter();if(order.join()!=='1,0,2')throw Error('sort');
fields['score-status'].value='failed';filter();if(nodes.filter(c=>!c.hidden)[0].id!==2)throw Error('failed filter');
fields['score-sort'].value='';filter();if(order.join()!=='0,1,2')throw Error('restore order');
"""
                    subprocess.run([node, "-"], input=harness + script + checks, text=True, encoding="utf-8", capture_output=True, check=True)
                record = load_journal(directory / "attributions.jsonl")[0]
                self.assertNotIn("score", record)
                await run_attribution("r", config(), with_score=True, force=True, formats=())
                self.assertEqual((score_api.await_count, attribute_api.await_count), (1, 2))
                row["human_note"] += "changed"
                source.write_text(json.dumps(row) + "\n", encoding="utf-8")
                await run_attribution("r", config(), with_score=True, formats=())
                self.assertEqual((score_api.await_count, attribute_api.await_count), (1, 3))
                await run_attribution("r", config(), with_score=True, report_only=True, formats=())
                self.assertEqual((score_api.await_count, attribute_api.await_count), (1, 3))
                attribute_api.side_effect = RuntimeError("attribute failed")
                await run_attribution("r", config(), with_score=True, force=True, formats=())
                self.assertEqual((score_api.await_count, attribute_api.await_count), (1, 4))
                attribute_api.side_effect = None
                await run_attribution("r", config(), with_score=True, formats=())
                self.assertEqual((score_api.await_count, attribute_api.await_count), (1, 5))
                loader.side_effect = ValueError("bad scoring config")
                await run_attribution("r", config(), with_score=True, force=True, formats=("html",))
                self.assertEqual(attribute_api.await_count, 6)
                self.assertIn("评分配置加载或执行失败", (directory / "report.html").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
