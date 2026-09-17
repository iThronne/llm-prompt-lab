"""从归因结构化数据生成 HTML 和 Excel；不调用模型、不读取评分。"""

import base64
import json
from pathlib import Path
import re

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.attribution_schema import STAGES, STATUSES

LABELS = {
    **STAGES, **STATUSES, "success": "执行成功", "failed": "执行失败", "not_processed": "未处理/输入已变化",
    "core": "核心影响", "important": "重要影响", "local": "局部影响",
    "direct": "直接支持", "indirect": "间接支持", "supported": "材料支持",
    "refuted": "材料反驳", "unverifiable": "无法验证", "checked": "已检查",
    "not_applicable": "不适用", "unavailable": "无法检查",
}


def label(value):
    return LABELS.get(value, value) if value is not None else "未定"


def anchor(case_id, path):
    return "source-" + base64.urlsafe_b64encode((case_id + "\0" + path).encode()).decode().rstrip("=")


def source_leaves(value, path=""):
    """保留全部文本值，JSON Pointer 可从证据跳转；完整 JSON 另行展示。"""
    if isinstance(value, str):
        yield {"path": path, "text": value}
    elif isinstance(value, dict):
        for key, item in value.items():
            escaped = key.replace("~", "~0").replace("/", "~1")
            yield from source_leaves(item, path + "/" + escaped)
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from source_leaves(item, path + f"/{index}")


def summary_rows(summary):
    yield ["总案例数", "全部当前案例", summary["total_cases"]]
    categories = {
        "execution_counts": "执行状态（案例数）", "status_counts": "归因状态（成功案例数）",
        "unique_primary_cases_by_stage": "唯一主因（案例数）", "joint_primary_cases_by_stage": "并列主因涉及（案例数）",
        "involved_cases_by_stage": "全部因素涉及（案例数，可重叠）",
        "problem_type_case_counts": "具体问题类型（案例数，可重叠）", "note_claim_counts": "评论核对（主张数）",
        "failure_counts": "执行失败原因（案例数）",
    }
    for key, title in categories.items():
        for name, count in summary.get(key, {}).items():
            yield [title, label(name), count]


def report_tables(records, summary):
    tables = {
        "案例总览": (["Case ID", "原始行号", "Query 原文", "Query 中文概要", "归因/执行状态", "主要因素", "因素数量", "中文归因摘要", "输入问题", "执行错误", "完整记录定位"], []),
        "回答问题": (["Case ID", "问题 ID", "问题描述", "影响程度", "证据 ID", "完整记录定位"], []),
        "因素明细": (["Case ID", "因素 ID", "优先级", "问题环节", "问题类型", "问题描述", "影响问题 ID", "影响程度", "证据强度", "证据 ID", "作用机制", "排序依据", "改进建议", "验证方法", "完整记录定位"], []),
        "证据明细": (["Case ID", "证据 ID", "来源", "原文路径", "原文证据", "原文语言", "中文释义", "翻译歧义说明", "完整记录定位"], []),
        "评论核对": (["Case ID", "评论主张", "核对结果", "判定依据", "证据 ID", "完整记录定位"], []),
        "建议与待确认": (["Case ID", "类型", "环节", "说明", "改进动作/所需材料", "验证方式/判断限制", "证据 ID", "完整记录定位"], []),
        "环节检查": (["Case ID", "环节", "检查状态", "发现", "证据 ID", "完整记录定位"], []),
        "汇总统计": (["统计口径", "项目", "数量"], list(summary_rows(summary))),
    }
    for record in records:
        cid = record["case_id"]
        locator = f"attributions.jsonl | case_id={cid} | input_hash={record['input_hash']} | 最新匹配记录"
        result = record.get("result") or {}
        factors = result.get("factors", [])
        snapshot = record["input_snapshot"]
        primary = "；".join(f["description"] for f in factors if f["id"] in result.get("primary_factor_ids", []))
        tables["案例总览"][1].append([cid, record["row_index"], snapshot["query"],
            result.get("content_summaries_zh", {}).get("query"),
            label(result.get("status") or record["execution"]["status"]), primary, len(factors),
            result.get("summary"), "\n".join(record["input_quality"]), record["execution"].get("error"), locator])
        for issue in result.get("issues", []):
            tables["回答问题"][1].append([cid, issue["id"], issue["description"], label(issue["impact_level"]),
                                      ", ".join(issue["evidence_ids"]), locator])
        for factor in factors:
            tables["因素明细"][1].append([cid, factor["id"], factor["priority"], label(factor["stage"]),
                factor["problem_type"], factor["description"], ", ".join(factor["issue_ids"]),
                label(factor["impact_level"]), label(factor["evidence_strength"]), ", ".join(factor["evidence_ids"]),
                factor["mechanism"], factor["ranking_reason"], factor["fix"]["action"], factor["fix"]["validation"], locator])
        for evidence in result.get("evidence", []):
            tables["证据明细"][1].append([cid, evidence["id"], label(evidence["source"]), evidence["path"], evidence["quote"],
                evidence["source_language"], evidence["translation_zh"], evidence["translation_note"], locator])
        for note in result.get("note_assessment", []):
            tables["评论核对"][1].append([cid, note["claim"], label(note["assessment"]), note["reason"],
                                      ", ".join(note["evidence_ids"]), locator])
        for suggestion in result.get("suggestions", []):
            tables["建议与待确认"][1].append([cid, "未观察到实际影响的建议", label(suggestion["stage"]),
                suggestion["description"], suggestion["fix"]["action"], suggestion["fix"]["validation"],
                ", ".join(suggestion["evidence_ids"]), locator])
        for question in result.get("open_questions", []):
            tables["建议与待确认"][1].append([cid, "待确认", None, question["question"],
                question["needed_material"], question["limits"], None, locator])
        for check in result.get("stage_checks", []):
            tables["环节检查"][1].append([cid, label(check["stage"]), label(check["status"]),
                                        check["finding"], ", ".join(check["evidence_ids"]), locator])
    return tables


def excel_text(value):
    """UTF-16 长度限制及 XML 控制字符显式摘录；JSONL 保留原文。"""
    if not isinstance(value, str):
        return value
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "�", value)
    note = ""
    if cleaned != value:
        note = "\n[Excel 不支持的控制字符已替换；原文见 JSONL]"
    encoded = cleaned.encode("utf-16-le")
    if len(encoded) > 60000:
        cleaned = encoded[:60000].decode("utf-16-le", errors="ignore")
        note += "\n[展示摘录：全文见本行完整记录定位对应的 JSONL]"
    return cleaned + note


def export_attribution_excel(path: Path, records, summary):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, (headers, rows) in report_tables(records, summary).items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            values = [excel_text(value) for value in row]
            sheet.append(values)
            for cell, value in zip(sheet[sheet.max_row], values):
                if isinstance(value, str):
                    cell.data_type = "s"  # 原文以 = 开头也不是可执行公式。
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[sheet.max_row].height = 75
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 30
        sheet.freeze_panes = "C2" if name != "汇总统计" else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, 1):
            width = 48 if any(word in header for word in ("描述", "证据", "说明", "摘要", "依据", "建议", "原文", "机制", "定位")) else 24
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.sheet_view.zoomScale = 85
    # 读取说明独立放置，避免改变可筛选表的第一行。
    notes = workbook.create_sheet("阅读说明", 0)
    for row in [["项目", "说明"], ["完整数据", "attributions.jsonl 保存完整输入和结构化结果，Excel 仅为阅读视图。"],
                ["长文本", "单元格标注展示摘录时，按完整记录定位查 JSONL。较长内容可在编辑栏阅读或查看 HTML。"],
                ["优先级", "1 为最高优先级；相同数字为并列；空白表示无法可靠排序。"],
                ["中文概要/释义", "仅辅助阅读，不能替代原文作为证据。"],
                ["统计口径", summary["counting_note"]],
                ["未处理案例", "未处理/输入已变化表示当前输入没有匹配结果，不使用过期归因。"]]:
        notes.append(row)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 95
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[row[0].row].height = 40
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(path)


def generate_reports(directory: Path, records, summary, formats=("html", "xlsx")):
    unknown = set(formats) - {"html", "xlsx"}
    if unknown:
        raise ValueError(f"未知报告格式: {unknown}")
    if "html" in formats:
        environment = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                                  autoescape=select_autoescape(["html"]))
        environment.globals.update(label=label, anchor=anchor, source_leaves=source_leaves,
                                   sorted_factors=lambda factors: sorted(factors, key=lambda f: (f["priority"] is None, f["priority"] or 0, f["id"])),
                                   full_json=lambda value: json.dumps(value, ensure_ascii=False, indent=2))
        rendered = environment.get_template("attribution.html").render(
            records=records, summary=summary, stats=list(summary_rows(summary)), stages=STAGES, statuses=STATUSES)
        temporary = directory / "report.html.tmp"
        temporary.write_text(rendered, encoding="utf-8")
        temporary.replace(directory / "report.html")
    if "xlsx" in formats:
        export_attribution_excel(directory / "report.xlsx", records, summary)
