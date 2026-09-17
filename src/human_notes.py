"""人工评论的 XLSX 匹配、版本校验与独立 JSONL 存储。"""

from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import uuid
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

from src.attribution import digest, read_responses

MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_XLSX_EXPANDED = 100 * 1024 * 1024
MAX_SHEET_ROWS = 50000
MAX_COLUMNS = 500


class NoteConflict(ValueError):
    def __init__(self, rows):
        super().__init__("这些案例的回答或评论已变化，请重新加载后核对，未保存任何修改。")
        self.rows = rows


def answer_text(row):
    try:
        value = row["response"]["choices"][0]["message"].get("content")
    except (KeyError, IndexError, TypeError, AttributeError):
        value = row.get("answer")
    if value is None:
        return ""
    return value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)


def normalize_query(value):
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def target_hash(row):
    # 相同 Query/Answer 但上下文已更换时，也不能沿用旧评论而不提示。
    return digest({"query": row.get("query"), "answer": answer_text(row),
                   "rendered_request": row.get("rendered_request")})


def load_notes(path: Path):
    records = []
    if path.exists():
        with path.open(encoding="utf-8") as stream:
            for line in stream:
                if line.strip():
                    record = json.loads(line)
                    if (type(record.get("row_index")) is not int
                            or not isinstance(record.get("human_note"), str)
                            or not isinstance(record.get("target_hash"), str)
                            or not isinstance(record.get("answer_hash"), str)
                            or not isinstance(record.get("revision"), str)):
                        raise ValueError("human_notes.jsonl 记录格式无效，未忽略损坏数据。")
                    records.append(record)
    return records


def apply_human_notes(rows, run_dir: Path):
    """应用最新且与当前 Query/Answer/messages 匹配的人工评论；不修改源记录。"""
    latest = {n["row_index"]: n for n in load_notes(run_dir / "human_notes.jsonl")}
    result = []
    for row in rows:
        copy = dict(row)
        note = latest.get(row["row_index"])
        if note and note["target_hash"] == target_hash(row):
            copy.update(human_note=note["human_note"], human_note_answer_hash=note["answer_hash"])
        elif note:
            copy["human_note_warning"] = "页面保存的评论属于旧版 Query/Answer/messages，未自动套用；请在人工评论页面核对。"
        result.append(copy)
    return result


def note_state(row, latest):
    note = latest.get(row["row_index"])
    current_target = target_hash(row)
    original = row.get("human_note")
    if original is None:
        original = row.get("note")
    original = "" if original is None else str(original)
    matches = note is not None and note["target_hash"] == current_target
    effective = note["human_note"] if matches else original
    revision = digest({"target": current_target, "base_note": original,
                       "base_binding": row.get("human_note_answer_hash"),
                       "saved_revision": note["revision"] if note else None})
    return {"row_index": row["row_index"], "query": str(row.get("query") or ""),
            "answer_preview": answer_text(row)[:500], "human_note": effective,
            "revision": revision, "has_note": bool(effective.strip()),
            "note_source": "页面保存" if matches else ("源数据" if original else "尚未填写"),
            "stale_note": note["human_note"] if note and not matches else None}


def list_cases(run_dir: Path):
    rows = read_responses(run_dir / "responses.jsonl")
    latest = {n["row_index"]: n for n in load_notes(run_dir / "human_notes.jsonl")}
    cases = [note_state(row, latest) for row in rows]
    counts = Counter(normalize_query(c["query"]) for c in cases)
    for case in cases:
        case["query_match_count"] = counts[normalize_query(case["query"])]
    return cases


@contextmanager
def notes_lock(run_dir):
    path = run_dir / ".human_notes.lock"
    try:
        descriptor = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise ValueError("另一个任务正在保存人工评论；若是异常退出遗留，请确认无保存任务后移除 .human_notes.lock。") from exc
    try:
        with os.fdopen(descriptor, "w") as stream:
            stream.write(str(os.getpid()))
        yield
    finally:
        path.unlink(missing_ok=True)


def save_notes(run_dir: Path, changes):
    """校验全部修改后原子替换评论日志；乐观锁保护其他页面的更新。"""
    if not isinstance(changes, list) or not changes:
        raise ValueError("changes 必须是非空数组")
    seen = set()
    for change in changes:
        if not isinstance(change, dict):
            raise ValueError("每项修改必须是对象")
        index = change.get("row_index")
        if type(index) is not int or index in seen:
            raise ValueError("row_index 必须是唯一整数")
        if not isinstance(change.get("human_note"), str) or not isinstance(change.get("revision"), str):
            raise ValueError("human_note、revision 必须为字符串")
        seen.add(index)
    with notes_lock(run_dir):
        source = run_dir / "responses.jsonl"
        source_hash = digest(source.read_text(encoding="utf-8-sig"))
        rows = {r["row_index"]: r for r in read_responses(source)}
        path = run_dir / "human_notes.jsonl"
        history = load_notes(path)
        latest = {n["row_index"]: n for n in history}
        conflicts = [c["row_index"] for c in changes if c["row_index"] not in rows
                     or note_state(rows[c["row_index"]], latest)["revision"] != c["revision"]]
        if conflicts:
            raise NoteConflict(conflicts)
        for change in changes:
            row = rows[change["row_index"]]
            history.append({"row_index": row["row_index"], "human_note": change["human_note"],
                            "target_hash": target_hash(row), "answer_hash": digest(answer_text(row)),
                            "revision": uuid.uuid4().hex, "saved_at": datetime.now(timezone.utc).isoformat()})
        temp = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
        try:
            with temp.open("w", encoding="utf-8") as stream:
                for record in history:
                    stream.write(json.dumps(record, ensure_ascii=False, allow_nan=False) + "\n")
                stream.flush()
                os.fsync(stream.fileno())
            if source_hash != digest(source.read_text(encoding="utf-8-sig")):
                raise NoteConflict([c["row_index"] for c in changes])
            temp.replace(path)
        finally:
            temp.unlink(missing_ok=True)
    return {"saved": len(changes)}


def open_xlsx(data):
    if not data or len(data) > MAX_XLSX_BYTES:
        raise ValueError("XLSX 为空或超过 20 MB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_EXPANDED:
                raise ValueError("XLSX 解压后超过 100 MB，请拆分文件")
        return load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("无法读取 XLSX，请上传有效的 .xlsx 文件") from exc


def workbook_info(data):
    workbook = open_xlsx(data)
    try:
        sheets = []
        for sheet in workbook:
            if (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("工作表列数超过 500，请只保留相关列")
            headers = next(sheet.iter_rows(min_row=1, max_row=1), ())
            sheets.append({"name": sheet.title, "columns": [
                {"index": i, "label": str(cell.value) if cell.value is not None else f"未命名列 {i + 1}"}
                for i, cell in enumerate(headers)], "rows": max(0, (sheet.max_row or 1) - 1)})
        return sheets
    finally:
        workbook.close()


def preview_import(data, cases, sheet_name, query_column, note_column):
    if type(query_column) is not int or type(note_column) is not int or query_column == note_column:
        raise ValueError("请选择不同的 Query 列与人工评论列")
    workbook = open_xlsx(data)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("工作表不存在")
        sheet = workbook[sheet_name]
        if min(query_column, note_column) < 0 or max(query_column, note_column) >= (sheet.max_column or 0):
            raise ValueError("列索引超出工作表范围")
        if (sheet.max_row or 0) > MAX_SHEET_ROWS + 1 or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ValueError("单表最多支持 50000 条数据、500 列，请拆分文件")
        imported = []
        for excel_row, cells in enumerate(sheet.iter_rows(min_row=2, max_col=max(query_column, note_column) + 1), 2):
            if excel_row > MAX_SHEET_ROWS + 1:
                raise ValueError("数据行数超过 50000")
            query_cell, note_cell = cells[query_column], cells[note_column]
            if query_cell.value is None and note_cell.value is None:
                continue
            query = "" if query_cell.value is None else str(query_cell.value)
            note = "" if note_cell.value is None else str(note_cell.value)
            warnings = []
            blocked = False
            if query_cell.data_type in {"f", "e"} or note_cell.data_type in {"f", "e"}:
                warnings.append("Query 或评论为公式/错误单元格，请先在 Excel 中转换为文本值")
                blocked = True
            if not normalize_query(query):
                warnings.append("Query 为空")
                blocked = True
            if not note.strip():
                warnings.append("评论为空，不导入以免清空已有评论；可在案例中手动清空")
                blocked = True
            if len(query) >= 32767 or len(note) >= 32767:
                warnings.append("单元格达到 Excel 长度上限，原始内容可能已截断，请核对")
            imported.append({"excel_row": excel_row, "query": query, "human_note": note,
                             "warnings": warnings, "blocked": blocked})
    finally:
        workbook.close()
    source_counts = Counter(normalize_query(r["query"]) for r in imported if normalize_query(r["query"]))
    targets = defaultdict(list)
    for case in cases:
        targets[normalize_query(case["query"])].append(case)
    stats = Counter()
    for row in imported:
        matches = targets.get(normalize_query(row["query"]), []) if normalize_query(row["query"]) else []
        duplicate_source = source_counts[normalize_query(row["query"])] > 1
        if duplicate_source:
            row["warnings"].append("表内 Query 重复：需选择采用哪一行评论")
            stats["duplicate_source_rows"] += 1
        if len(matches) > 1:
            row["warnings"].append("目标 Query 重复：请结合回答选择对应案例，不自动批量填充")
            stats["duplicate_target_rows"] += 1
        if not matches:
            row["warnings"].append("未匹配到当前实验的 Query")
            stats["unmatched_rows"] += 1
        if any(c["has_note"] for c in matches):
            row["warnings"].append("匹配案例已有评论：应用后会替换，请核对")
            stats["existing_note_rows"] += 1
        if row["blocked"]:
            stats["blocked_rows"] += 1
        row["candidates"] = [{"row_index": c["row_index"], "answer_preview": c["answer_preview"],
                               "has_note": c["has_note"], "revision": c["revision"]} for c in matches]
        row["auto_select"] = (not row["blocked"] and len(matches) == 1 and not row["warnings"])
        if row["auto_select"]:
            stats["ready_rows"] += 1
    return {"rows": imported, "summary": {"total_rows": len(imported), **dict(stats)},
            "matching_rule": "Query 统一换行并去掉首尾空白后精确匹配；不忽略大小写，不做模糊匹配。"}
